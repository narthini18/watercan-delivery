from flask import Flask, render_template, request, jsonify, session, redirect, url_for, send_file
from datetime import datetime, timedelta
import sqlite3, os, io, hashlib
from functools import wraps
try:
    import qrcode
    from PIL import Image
    HAS_QR = True
except ImportError:
    HAS_QR = False

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'watercan-super-secret-2024-change-in-prod')

DB_PATH = os.path.join(app.instance_path, 'watercan.db')
PRICE_PER_CAN = int(os.environ.get('PRICE_PER_CAN', 50))

# ─── DB INIT ──────────────────────────────────────────────────────────────────

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA journal_mode=WAL")
    return conn

def init_db():
    os.makedirs(app.instance_path, exist_ok=True)
    with get_db() as conn:
        conn.executescript('''
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                password_hash TEXT NOT NULL,
                role TEXT NOT NULL CHECK(role IN ('customer','worker','owner')),
                name TEXT NOT NULL,
                flat_no TEXT DEFAULT '',
                phone TEXT DEFAULT '',
                address TEXT DEFAULT '',
                created_at TEXT DEFAULT (datetime('now'))
            );

            CREATE TABLE IF NOT EXISTS orders (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                order_id TEXT UNIQUE NOT NULL,
                username TEXT NOT NULL,
                flat_no TEXT,
                quantity INTEGER NOT NULL,
                order_date TEXT NOT NULL,
                order_time TEXT NOT NULL,
                status TEXT DEFAULT 'pending' CHECK(status IN ('pending','delivered','cancelled')),
                delivered_date TEXT,
                delivered_time TEXT,
                delivered_by TEXT,
                amount REAL NOT NULL,
                payment_status TEXT DEFAULT 'pending' CHECK(payment_status IN ('pending','paid')),
                FOREIGN KEY(username) REFERENCES users(username)
            );

            CREATE TABLE IF NOT EXISTS payments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                payment_id TEXT UNIQUE NOT NULL,
                username TEXT NOT NULL,
                customer_name TEXT,
                flat_no TEXT,
                order_id TEXT NOT NULL,
                amount REAL NOT NULL,
                payment_date TEXT NOT NULL,
                payment_time TEXT NOT NULL,
                payment_method TEXT DEFAULT 'cash',
                notes TEXT DEFAULT '',
                recorded_by TEXT,
                FOREIGN KEY(username) REFERENCES users(username)
            );

            CREATE TABLE IF NOT EXISTS notifications (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT NOT NULL,
                message TEXT NOT NULL,
                type TEXT DEFAULT 'info',
                read INTEGER DEFAULT 0,
                created_at TEXT DEFAULT (datetime('now'))
            );
        ''')
        # Seed default users if empty
        row = conn.execute("SELECT COUNT(*) as c FROM users").fetchone()
        if row['c'] == 0:
            seed_users = [
                ('customer1', hash_pw('pass123'), 'customer', 'John Doe',    'A1-203', '9876543210', 'Building A, Flat 203'),
                ('customer2', hash_pw('pass123'), 'customer', 'Priya Sharma','B2-101', '9876543213', 'Building B, Flat 101'),
                ('worker1',   hash_pw('pass123'), 'worker',   'Raju Kumar',  '',       '9876543211', ''),
                ('owner',     hash_pw('admin123'),'owner',    'Owner Admin',  '',       '9876543212', ''),
            ]
            conn.executemany(
                "INSERT INTO users(username,password_hash,role,name,flat_no,phone,address) VALUES(?,?,?,?,?,?,?)",
                seed_users
            )
        conn.commit()

def hash_pw(pw):
    return hashlib.sha256(pw.encode()).hexdigest()

# ─── AUTH ─────────────────────────────────────────────────────────────────────

def login_required(f):
    @wraps(f)
    def decorated(*args, **kwargs):
        if 'username' not in session:
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated

def role_required(*roles):
    def decorator(f):
        @wraps(f)
        def decorated(*args, **kwargs):
            if session.get('role') not in roles:
                return redirect(url_for('login'))
            return f(*args, **kwargs)
        return decorated
    return decorator

def add_notification(username, message, ntype='info'):
    try:
        with get_db() as conn:
            conn.execute(
                "INSERT INTO notifications(username,message,type) VALUES(?,?,?)",
                (username, message, ntype)
            )
            conn.commit()
    except Exception:
        pass

# ─── ROUTES ───────────────────────────────────────────────────────────────────

@app.route('/')
def index():
    if 'username' in session:
        role = session.get('role')
        return redirect(url_for(f'{role}_dashboard'))
    return redirect(url_for('login'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        data = request.get_json()
        username = data.get('username', '').strip()
        password = data.get('password', '')
        with get_db() as conn:
            user = conn.execute(
                "SELECT * FROM users WHERE username=? AND password_hash=?",
                (username, hash_pw(password))
            ).fetchone()
        if user:
            session.update({
                'username': user['username'],
                'role': user['role'],
                'name': user['name'],
                'flat_no': user['flat_no'] or ''
            })
            return jsonify({'success': True, 'role': user['role']})
        return jsonify({'success': False, 'message': 'Invalid username or password'})
    return render_template('login.html')

@app.route('/logout')
def logout():
    session.clear()
    return redirect(url_for('login'))

# ─── CUSTOMER ─────────────────────────────────────────────────────────────────

@app.route('/customer/dashboard')
@login_required
@role_required('customer')
def customer_dashboard():
    return render_template('customer_dashboard.html')

@app.route('/customer/data')
@login_required
@role_required('customer')
def customer_data():
    username = session['username']
    with get_db() as conn:
        orders = [dict(r) for r in conn.execute(
            "SELECT * FROM orders WHERE username=? ORDER BY order_date DESC, order_time DESC LIMIT 50",
            (username,)
        ).fetchall()]
        payments = [dict(r) for r in conn.execute(
            "SELECT * FROM payments WHERE username=? ORDER BY payment_date DESC LIMIT 20",
            (username,)
        ).fetchall()]
        notifs = [dict(r) for r in conn.execute(
            "SELECT * FROM notifications WHERE username=? AND read=0 ORDER BY created_at DESC",
            (username,)
        ).fetchall()]

    total_paid   = sum(o['amount'] for o in orders if o['payment_status'] == 'paid')
    total_pend   = sum(o['amount'] for o in orders if o['payment_status'] == 'pending')
    total_cans   = sum(o['quantity'] for o in orders)
    pending_orders_count = sum(1 for o in orders if o['status'] == 'pending')

    monthly = {}
    for o in orders:
        m = o['order_date'][:7]
        if m not in monthly:
            monthly[m] = {'month': m, 'cans': 0, 'amount': 0, 'paid': 0}
        monthly[m]['cans']   += o['quantity']
        monthly[m]['amount'] += o['amount']
        if o['payment_status'] == 'paid':
            monthly[m]['paid'] += o['amount']

    return jsonify({
        'total_paid': total_paid,
        'total_pending': total_pend,
        'total_cans': total_cans,
        'pending_orders': pending_orders_count,
        'orders': orders,
        'payments': payments,
        'monthly': sorted(monthly.values(), key=lambda x: x['month'], reverse=True),
        'notifications': notifs
    })

@app.route('/customer/place-order', methods=['POST'])
@login_required
@role_required('customer')
def place_order():
    data = request.get_json()
    qty = max(1, min(20, int(data.get('quantity', 1))))
    now = datetime.now()
    order_id = f"ORD{now.strftime('%Y%m%d%H%M%S')}{session['username'][:3].upper()}"
    amount = qty * PRICE_PER_CAN
    with get_db() as conn:
        conn.execute(
            """INSERT INTO orders(order_id,username,flat_no,quantity,order_date,order_time,amount)
               VALUES(?,?,?,?,?,?,?)""",
            (order_id, session['username'], session['flat_no'], qty,
             now.strftime('%Y-%m-%d'), now.strftime('%H:%M:%S'), amount)
        )
        conn.commit()
    add_notification(session['username'], f"Order {order_id} placed — {qty} can(s) for ₹{amount}", 'order')
    return jsonify({'success': True, 'order_id': order_id, 'amount': amount})

@app.route('/customer/mark-notifications-read', methods=['POST'])
@login_required
def mark_notifs_read():
    with get_db() as conn:
        conn.execute("UPDATE notifications SET read=1 WHERE username=?", (session['username'],))
        conn.commit()
    return jsonify({'success': True})

# ─── WORKER ───────────────────────────────────────────────────────────────────

@app.route('/worker/dashboard')
@login_required
@role_required('worker')
def worker_dashboard():
    return render_template('worker_dashboard.html')

@app.route('/worker/data')
@login_required
@role_required('worker')
def worker_data():
    today = datetime.now().strftime('%Y-%m-%d')
    with get_db() as conn:
        pending = [dict(r) for r in conn.execute(
            "SELECT * FROM orders WHERE status='pending' ORDER BY order_date ASC, order_time ASC"
        ).fetchall()]
        my_deliveries = [dict(r) for r in conn.execute(
            "SELECT * FROM orders WHERE delivered_by=? AND status='delivered'",
            (session['username'],)
        ).fetchall()]
        today_count = sum(o['quantity'] for o in my_deliveries if o['delivered_date'] == today)
        total_count = sum(o['quantity'] for o in my_deliveries)

    monthly = {}
    for o in my_deliveries:
        m = o['delivered_date'][:7]
        monthly[m] = monthly.get(m, 0) + o['quantity']

    return jsonify({
        'pending_orders': pending,
        'today_cans': today_count,
        'total_delivered': total_count,
        'monthly': [{'month': k, 'cans': v} for k, v in sorted(monthly.items(), reverse=True)]
    })

@app.route('/worker/mark-delivered', methods=['POST'])
@login_required
@role_required('worker')
def mark_delivered():
    data = request.get_json()
    order_id = data.get('order_id')
    now = datetime.now()
    with get_db() as conn:
        order = conn.execute("SELECT * FROM orders WHERE order_id=?", (order_id,)).fetchone()
        if not order:
            return jsonify({'success': False})
        conn.execute(
            """UPDATE orders SET status='delivered', delivered_date=?, delivered_time=?, delivered_by=?
               WHERE order_id=?""",
            (now.strftime('%Y-%m-%d'), now.strftime('%H:%M:%S'), session['username'], order_id)
        )
        conn.commit()
    add_notification(order['username'],
        f"Your order {order_id} has been delivered! 💧", 'delivery')
    return jsonify({'success': True})

# ─── OWNER ────────────────────────────────────────────────────────────────────

@app.route('/owner/dashboard')
@login_required
@role_required('owner')
def owner_dashboard():
    return render_template('owner_dashboard.html')

@app.route('/owner/data')
@login_required
@role_required('owner')
def owner_data():
    with get_db() as conn:
        orders  = [dict(r) for r in conn.execute("SELECT * FROM orders ORDER BY order_date DESC").fetchall()]
        payments = [dict(r) for r in conn.execute("SELECT * FROM payments ORDER BY payment_date DESC LIMIT 50").fetchall()]
        users   = [dict(r) for r in conn.execute("SELECT * FROM users WHERE role='customer'").fetchall()]
        workers = [dict(r) for r in conn.execute("SELECT * FROM users WHERE role='worker'").fetchall()]

    total_orders  = len(orders)
    total_cans    = sum(o['quantity'] for o in orders)
    total_revenue = sum(o['amount'] for o in orders if o['payment_status'] == 'paid')
    pending_amount = sum(o['amount'] for o in orders if o['payment_status'] == 'pending')

    # daily last 7 days
    daily = []
    for i in range(6, -1, -1):
        d = (datetime.now() - timedelta(days=i)).strftime('%Y-%m-%d')
        cans = sum(o['quantity'] for o in orders if o['order_date'] == d)
        daily.append({'date': d, 'cans': cans})

    # monthly
    monthly = {}
    for o in orders:
        m = o['order_date'][:7]
        if m not in monthly:
            monthly[m] = {'month': m, 'cans': 0, 'revenue': 0, 'pending': 0}
        monthly[m]['cans'] += o['quantity']
        if o['payment_status'] == 'paid':
            monthly[m]['revenue'] += o['amount']
        else:
            monthly[m]['pending'] += o['amount']

    # customer stats
    customer_stats = []
    for u in users:
        co = [o for o in orders if o['username'] == u['username']]
        customer_stats.append({
            'username': u['username'],
            'name': u['name'],
            'flat_no': u['flat_no'],
            'phone': u['phone'],
            'total_cans': sum(o['quantity'] for o in co),
            'paid': sum(o['amount'] for o in co if o['payment_status'] == 'paid'),
            'pending': sum(o['amount'] for o in co if o['payment_status'] == 'pending'),
        })

    # worker stats
    worker_stats = []
    for w in workers:
        wd = [o for o in orders if o['delivered_by'] == w['username'] and o['status'] == 'delivered']
        today = datetime.now().strftime('%Y-%m-%d')
        worker_stats.append({
            'name': w['name'],
            'username': w['username'],
            'today': sum(o['quantity'] for o in wd if o['delivered_date'] == today),
            'total': sum(o['quantity'] for o in wd)
        })

    # pending payments list
    pending_pay = [dict(o) for o in orders if o['payment_status'] == 'pending']
    for o in pending_pay:
        cust = next((u for u in users if u['username'] == o['username']), None)
        o['customer_name'] = cust['name'] if cust else o['username']

    return jsonify({
        'total_orders': total_orders,
        'total_cans': total_cans,
        'total_revenue': total_revenue,
        'pending_amount': pending_amount,
        'daily': daily,
        'monthly': sorted(monthly.values(), key=lambda x: x['month'], reverse=True),
        'customer_stats': customer_stats,
        'worker_stats': worker_stats,
        'pending_payment_orders': sorted(pending_pay, key=lambda x: x['order_date']),
        'payment_records': payments
    })

@app.route('/owner/record-payment', methods=['POST'])
@login_required
@role_required('owner')
def record_payment():
    data = request.get_json()
    order_id = data.get('order_id')
    method   = data.get('payment_method', 'cash')
    notes    = data.get('notes', '')
    now = datetime.now()
    with get_db() as conn:
        order = conn.execute("SELECT o.*, u.name as cname FROM orders o JOIN users u ON o.username=u.username WHERE o.order_id=?", (order_id,)).fetchone()
        if not order:
            return jsonify({'success': False, 'message': 'Order not found'})
        pay_id = f"PAY{now.strftime('%Y%m%d%H%M%S')}"
        conn.execute(
            """INSERT INTO payments(payment_id,username,customer_name,flat_no,order_id,amount,payment_date,payment_time,payment_method,notes,recorded_by)
               VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
            (pay_id, order['username'], order['cname'], order['flat_no'], order_id,
             order['amount'], now.strftime('%Y-%m-%d'), now.strftime('%H:%M:%S'),
             method, notes, session['username'])
        )
        conn.execute("UPDATE orders SET payment_status='paid' WHERE order_id=?", (order_id,))
        conn.commit()
    add_notification(order['username'],
        f"Payment of ₹{order['amount']} received for {order_id} ({method}). Thank you! 🙏", 'payment')
    return jsonify({'success': True})

@app.route('/owner/record-payment-bulk', methods=['POST'])
@login_required
@role_required('owner')
def record_payment_bulk():
    data     = request.get_json()
    username = data.get('username')
    method   = data.get('payment_method', 'cash')
    notes    = data.get('notes', '')
    now = datetime.now()
    with get_db() as conn:
        pending = conn.execute(
            "SELECT o.*, u.name as cname FROM orders o JOIN users u ON o.username=u.username WHERE o.username=? AND o.payment_status='pending'",
            (username,)
        ).fetchall()
        if not pending:
            return jsonify({'success': False, 'message': 'No pending payments'})
        total = 0
        for i, o in enumerate(pending):
            pay_id = f"PAY{now.strftime('%Y%m%d%H%M%S')}{i}"
            conn.execute(
                """INSERT INTO payments(payment_id,username,customer_name,flat_no,order_id,amount,payment_date,payment_time,payment_method,notes,recorded_by)
                   VALUES(?,?,?,?,?,?,?,?,?,?,?)""",
                (pay_id, o['username'], o['cname'], o['flat_no'], o['order_id'],
                 o['amount'], now.strftime('%Y-%m-%d'), now.strftime('%H:%M:%S'),
                 method, notes, session['username'])
            )
            total += o['amount']
        conn.execute("UPDATE orders SET payment_status='paid' WHERE username=? AND payment_status='pending'", (username,))
        conn.commit()
    add_notification(username,
        f"Bulk payment of ₹{total:.0f} received for {len(pending)} orders ({method}). Thank you! 🙏", 'payment')
    return jsonify({'success': True, 'total': total, 'count': len(pending)})

@app.route('/owner/add-customer', methods=['POST'])
@login_required
@role_required('owner')
def add_customer():
    data = request.get_json()
    username = data.get('username', '').strip().lower()
    name     = data.get('name', '').strip()
    flat_no  = data.get('flat_no', '').strip()
    phone    = data.get('phone', '').strip()
    password = data.get('password', 'pass123')
    if not username or not name:
        return jsonify({'success': False, 'message': 'Username and name are required'})
    try:
        with get_db() as conn:
            conn.execute(
                "INSERT INTO users(username,password_hash,role,name,flat_no,phone) VALUES(?,?,?,?,?,?)",
                (username, hash_pw(password), 'customer', name, flat_no, phone)
            )
            conn.commit()
        return jsonify({'success': True})
    except sqlite3.IntegrityError:
        return jsonify({'success': False, 'message': 'Username already exists'})

@app.route('/owner/export-excel')
@login_required
@role_required('owner')
def export_excel():
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return "openpyxl not installed", 500

    with get_db() as conn:
        orders   = [dict(r) for r in conn.execute("SELECT * FROM orders ORDER BY order_date DESC").fetchall()]
        payments = [dict(r) for r in conn.execute("SELECT * FROM payments ORDER BY payment_date DESC").fetchall()]
        users    = [dict(r) for r in conn.execute("SELECT * FROM users WHERE role='customer'").fetchall()]

    wb = openpyxl.Workbook()
    hdr_fill = PatternFill("solid", fgColor="0EA5E9")
    hdr_font = Font(bold=True, color="FFFFFF")

    def make_sheet(wb, title, headers, rows, is_first=False):
        ws = wb.active if is_first else wb.create_sheet(title)
        ws.title = title
        ws.append(headers)
        for cell in ws[1]:
            cell.fill = hdr_fill
            cell.font = hdr_font
            cell.alignment = Alignment(horizontal='center')
        for row in rows:
            ws.append(row)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(len(str(col[0].value or '')), 12) + 2
        return ws

    # Sheet 1: Payments
    make_sheet(wb, 'Payments', ['Payment ID','Customer','Flat','Order ID','Amount','Date','Time','Method','Notes'],
        [[p['payment_id'],p['customer_name'],p['flat_no'],p['order_id'],f"₹{p['amount']:.0f}",
          p['payment_date'],p['payment_time'],p['payment_method'],p['notes']] for p in payments], True)

    # Sheet 2: Pending
    pending = [o for o in orders if o['payment_status'] == 'pending']
    make_sheet(wb, 'Pending Payments', ['Order ID','Username','Flat','Qty','Amount','Order Date','Delivery Status'],
        [[o['order_id'],o['username'],o['flat_no'],o['quantity'],f"₹{o['amount']:.0f}",
          o['order_date'],o['status']] for o in pending])

    # Sheet 3: Customer Summary
    rows = []
    for u in users:
        co = [o for o in orders if o['username'] == u['username']]
        rows.append([u['name'], u['flat_no'], u['phone'],
                     sum(o['quantity'] for o in co),
                     f"₹{sum(o['amount'] for o in co):.0f}",
                     f"₹{sum(o['amount'] for o in co if o['payment_status']=='paid'):.0f}",
                     f"₹{sum(o['amount'] for o in co if o['payment_status']=='pending'):.0f}"])
    make_sheet(wb, 'Customer Summary', ['Name','Flat','Phone','Total Cans','Total Amount','Paid','Pending'], rows)

    # Sheet 4: All Orders
    make_sheet(wb, 'All Orders', ['Order ID','Customer','Flat','Qty','Amount','Date','Status','Delivered By','Payment'],
        [[o['order_id'],o['username'],o['flat_no'],o['quantity'],f"₹{o['amount']:.0f}",
          o['order_date'],o['status'],o['delivered_by'] or '',o['payment_status']] for o in orders])

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                     as_attachment=True, download_name=f"watercan_report_{datetime.now().strftime('%Y%m%d')}.xlsx")

@app.route('/owner/qr-code')
@login_required
@role_required('owner')
def generate_qr():
    base_url = request.host_url.rstrip('/')
    if not HAS_QR:
        # Return a simple placeholder PNG
        from PIL import Image, ImageDraw
        img = Image.new('RGB', (200, 200), 'white')
        d = ImageDraw.Draw(img)
        d.text((10, 90), f"QR: {base_url}", fill='black')
        buf = io.BytesIO()
        img.save(buf, format='PNG')
        buf.seek(0)
        return send_file(buf, mimetype='image/png', download_name='watercan_qr.png')
    qr = qrcode.QRCode(version=1, box_size=10, border=4)
    qr.add_data(base_url)
    qr.make(fit=True)
    img = qr.make_image(fill_color="#0EA5E9", back_color="white")
    buf = io.BytesIO()
    img.save(buf, format='PNG')
    buf.seek(0)
    return send_file(buf, mimetype='image/png', download_name='watercan_qr.png')

# ─── PWA ──────────────────────────────────────────────────────────────────────

@app.route('/manifest.json')
def manifest():
    return jsonify({
        "name": "Water Can Delivery",
        "short_name": "WaterCan",
        "description": "Order and track water can deliveries",
        "start_url": "/",
        "display": "standalone",
        "background_color": "#0EA5E9",
        "theme_color": "#0EA5E9",
        "orientation": "portrait",
        "icons": [
            {"src": "/static/icons/icon-192.png", "sizes": "192x192", "type": "image/png"},
            {"src": "/static/icons/icon-512.png", "sizes": "512x512", "type": "image/png"}
        ]
    })

@app.route('/sw.js')
def service_worker():
    sw_content = """
const CACHE = 'watercan-v1';
const ASSETS = ['/', '/login', '/static/css/style.css'];
self.addEventListener('install', e => e.waitUntil(caches.open(CACHE).then(c => c.addAll(ASSETS))));
self.addEventListener('fetch', e => {
  if (e.request.method !== 'GET') return;
  e.respondWith(fetch(e.request).catch(() => caches.match(e.request)));
});
"""
    from flask import Response
    return Response(sw_content, mimetype='application/javascript')

# ─── ENTRY ────────────────────────────────────────────────────────────────────

if __name__ == '__main__':
    init_db()
    port = int(os.environ.get('PORT', 5000))
    debug = os.environ.get('FLASK_ENV', 'production') != 'production'
    app.run(host='0.0.0.0', port=port, debug=debug)
